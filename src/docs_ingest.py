"""Ingestão de documentos brutos: .txt, .docx, .doc (via textutil no macOS)."""
from __future__ import annotations

import subprocess
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
INPUTS = ROOT / "inputs"

DOC_GLOBS = ("*.txt", "*.docx", "*.doc", "*.md")


def _read_txt(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_docx(path: Path) -> str:
    from docx import Document  # type: ignore

    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_doc_legacy(path: Path) -> str:
    """Converte .doc binário via textutil (macOS) ou antiword se existir."""
    # macOS
    try:
        r = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(path)],
            check=True,
            capture_output=True,
            text=True,
            timeout=60,
        )
        if r.stdout.strip():
            return r.stdout
    except (FileNotFoundError, subprocess.CalledProcessError, subprocess.TimeoutExpired):
        pass

    try:
        r = subprocess.run(
            ["antiword", str(path)],
            check=True,
            capture_output=True,
            text=True,
            timeout=60,
        )
        return r.stdout
    except (FileNotFoundError, subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
        raise RuntimeError(
            f"Não foi possível ler {path.name}. Use .docx/.txt ou instale textutil/antiword."
        ) from e


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return _read_txt(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".doc":
        return _read_doc_legacy(path)
    raise ValueError(f"extensão não suportada: {suffix}")


def list_doc_files(directory: Path = INPUTS) -> list[Path]:
    files: list[Path] = []
    for pattern in DOC_GLOBS:
        files.extend(directory.glob(pattern))
    # ignora README de instrução
    return sorted(
        p for p in files if p.name.lower() not in {"readme.txt", "readme.md"} and p.is_file()
    )


def load_documents(directory: Path = INPUTS) -> list[dict[str, Any]]:
    docs: list[dict[str, Any]] = []
    for path in list_doc_files(directory):
        text = extract_text(path)
        docs.append(
            {
                "name": path.name,
                "path": str(path),
                "ext": path.suffix.lower(),
                "lines": text.count("\n") + (1 if text else 0),
                "chars": len(text),
                "est_tokens_raw": max(1, len(text) // 4) if text else 0,
                "text": text,
            }
        )
    return docs
